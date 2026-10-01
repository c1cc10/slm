#!/usr/bin/env python3
"""
tools/convert_to_text.py — Convertitore universale per corpus SLM

Formati supportati: PDF, DOCX, PPTX, EPUB, HTML, ODT, TXT, MD

Dipendenze:
    pip install pdfminer.six python-docx python-pptx odfpy ebooklib beautifulsoup4

Esempi:
    python3 tools/convert_to_text.py documento.pdf --out data/output.txt
    python3 tools/convert_to_text.py --dir ~/Libri/Gutenberg --out data/gutenberg_it.txt
    python3 tools/convert_to_text.py --dir ~/Documenti --ext pdf,docx --out data/tecnico.txt
    python3 tools/convert_to_text.py manuale.pdf --out data/corpus.txt --append
"""

import argparse
import os
import sys
from pathlib import Path


# ─────────────────────────────────────────────────────────────────────────────
# PULIZIA TESTO — stessa pipeline di epub_to_text.py
# ─────────────────────────────────────────────────────────────────────────────

def clean_text(text: str) -> str:
    """Normalizza caratteri speciali e whitespace."""
    text = text.replace(' ', ' ')   # narrow no-break space → spazio
    text = text.replace('\xa0', ' ')     # non-breaking space → spazio
    text = text.replace('﻿', '')    # BOM
    text = text.replace('©', '')        # copyright
    text = text.replace('…', '...')     # ellipsis → tre punti ASCII
    text = text.replace('�', '')   # replacement character
    text = text.replace('\r\n', '\n').replace('\r', '\n')

    # Collassa spazi multipli su stessa riga (preserva \n)
    lines = []
    for line in text.split('\n'):
        lines.append(' '.join(line.split()))
    text = '\n'.join(lines)

    # Collassa righe vuote multiple in una sola
    while '\n\n\n' in text:
        text = text.replace('\n\n\n', '\n\n')

    return text.strip()


# ─────────────────────────────────────────────────────────────────────────────
# EXTRACTORS — uno per formato
# ─────────────────────────────────────────────────────────────────────────────

def extract_txt(path: Path) -> str:
    try:
        return path.read_text(encoding='utf-8', errors='replace')
    except Exception as e:
        raise RuntimeError(f"TXT: {e}")


def extract_pdf(path: Path) -> str:
    try:
        from pdfminer.high_level import extract_text as _extract
    except ImportError:
        raise RuntimeError("PDF: installa pdfminer.six  →  pip install pdfminer.six")
    try:
        text = _extract(str(path))
        return text or ''
    except Exception as e:
        raise RuntimeError(f"PDF: {e}")


def extract_docx(path: Path) -> str:
    try:
        import docx as _docx
    except ImportError:
        raise RuntimeError("DOCX: installa python-docx  →  pip install python-docx")
    try:
        doc = _docx.Document(str(path))
        return '\n'.join(p.text for p in doc.paragraphs if p.text.strip())
    except Exception as e:
        raise RuntimeError(f"DOCX: {e}")


def extract_pptx(path: Path) -> str:
    try:
        from pptx import Presentation as _Prs
    except ImportError:
        raise RuntimeError("PPTX: installa python-pptx  →  pip install python-pptx")
    try:
        prs = _Prs(str(path))
        parts = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, 'text') and shape.text.strip():
                    parts.append(shape.text)
        return '\n'.join(parts)
    except Exception as e:
        raise RuntimeError(f"PPTX: {e}")


def extract_epub(path: Path) -> str:
    try:
        import ebooklib
        from ebooklib import epub as _epub
        from bs4 import BeautifulSoup
    except ImportError:
        raise RuntimeError("EPUB: installa ebooklib beautifulsoup4  →  pip install ebooklib beautifulsoup4")
    try:
        book  = _epub.read_epub(str(path), options={'ignore_ncx': True})
        parts = []
        for item in book.get_items_of_type(ebooklib.ITEM_DOCUMENT):
            soup = BeautifulSoup(item.get_content(), 'html.parser')
            for tag in soup(['script', 'style', 'nav']):
                tag.decompose()
            text = soup.get_text(separator='\n')
            if text.strip():
                parts.append(text)
        return '\n\n'.join(parts)
    except Exception as e:
        raise RuntimeError(f"EPUB: {e}")


def extract_html(path: Path) -> str:
    try:
        from bs4 import BeautifulSoup
    except ImportError:
        raise RuntimeError("HTML: installa beautifulsoup4  →  pip install beautifulsoup4")
    try:
        raw  = path.read_bytes()
        soup = BeautifulSoup(raw, 'html.parser')
        for tag in soup(['script', 'style', 'nav', 'header', 'footer', 'aside']):
            tag.decompose()
        return soup.get_text(separator='\n')
    except Exception as e:
        raise RuntimeError(f"HTML: {e}")


def extract_odt(path: Path) -> str:
    try:
        from odf import text as _odftext
        from odf.opendocument import load as _load
        from odf.element import Element
    except ImportError:
        raise RuntimeError("ODT: installa odfpy  →  pip install odfpy")
    try:
        doc   = _load(str(path))
        parts = []
        for para in doc.text.getElementsByType(_odftext.P):
            t = ''.join(n.data for n in para.childNodes
                        if hasattr(n, 'data'))
            if t.strip():
                parts.append(t)
        return '\n'.join(parts)
    except Exception as e:
        raise RuntimeError(f"ODT: {e}")


# ─────────────────────────────────────────────────────────────────────────────
# DISPATCHER
# ─────────────────────────────────────────────────────────────────────────────

_EXTRACTORS = {
    '.txt':  extract_txt,
    '.md':   extract_txt,
    '.pdf':  extract_pdf,
    '.docx': extract_docx,
    '.doc':  extract_docx,
    '.pptx': extract_pptx,
    '.ppt':  extract_pptx,
    '.epub': extract_epub,
    '.html': extract_html,
    '.htm':  extract_html,
    '.odt':  extract_odt,
}

def extract_any(path: Path) -> str:
    ext = path.suffix.lower()
    fn  = _EXTRACTORS.get(ext)
    if fn is None:
        raise RuntimeError(f"Formato non supportato: {ext}")
    return fn(path)


# ─────────────────────────────────────────────────────────────────────────────
# MAIN
# ─────────────────────────────────────────────────────────────────────────────

def convert_files(paths, out_path: Path, append: bool, min_chars: int):
    mode = 'a' if append else 'w'
    total_chars = 0
    ok, skipped = 0, 0

    with open(out_path, mode, encoding='utf-8') as f_out:
        for p in paths:
            try:
                raw  = extract_any(p)
                text = clean_text(raw)
                if len(text) < min_chars:
                    skipped += 1
                    print(f"  SKIP  {p.name}  ({len(text)} char < {min_chars})")
                    continue
                f_out.write(text)
                f_out.write('\n\n')
                total_chars += len(text)
                ok += 1
                print(f"  OK    {p.name}  ({len(text):,} char)")
            except RuntimeError as e:
                skipped += 1
                print(f"  ERR   {p.name}  {e}")

    mb = total_chars / 1_000_000
    print(f"\n  Convertiti : {ok}  saltati : {skipped}")
    print(f"  Output     : {out_path}  ({mb:.2f} MB di testo)")


def collect_paths(args) -> list:
    """Raccoglie tutti i file da convertire in base agli argomenti."""
    exts = None
    if args.ext:
        exts = {('.' + e.strip().lstrip('.').lower()) for e in args.ext.split(',')}

    paths = []

    if args.dir:
        root = Path(args.dir)
        if not root.is_dir():
            print(f"  ERRORE: directory non trovata: {root}")
            sys.exit(1)
        for p in sorted(root.rglob('*')):
            if not p.is_file():
                continue
            if exts:
                if p.suffix.lower() not in exts:
                    continue
            elif p.suffix.lower() not in _EXTRACTORS:
                continue
            paths.append(p)

    if args.inputs:
        for s in args.inputs:
            p = Path(s)
            if not p.is_file():
                print(f"  AVVISO: file non trovato: {p}")
                continue
            paths.append(p)

    return paths


def main():
    parser = argparse.ArgumentParser(
        description='Convertitore universale → testo per corpus SLM',
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""esempi:
  python3 tools/convert_to_text.py documento.pdf --out data/out.txt
  python3 tools/convert_to_text.py --dir ~/Libri --out data/gutenberg_it.txt
  python3 tools/convert_to_text.py --dir ~/Docs --ext pdf,docx --out data/tecnico.txt
  python3 tools/convert_to_text.py slide.pptx --out data/corpus.txt --append
""")
    parser.add_argument('inputs',   nargs='*',
                        help='file da convertire (pdf, docx, pptx, epub, html, odt, txt)')
    parser.add_argument('--dir',    default=None,
                        help='directory da scansionare ricorsivamente')
    parser.add_argument('--ext',    default=None,
                        help='filtro estensioni: pdf,docx,html (default: tutti i formati)')
    parser.add_argument('--out',    required=True,
                        help='file di output .txt')
    parser.add_argument('--append', action='store_true',
                        help='aggiunge al file esistente (default: sovrascrive)')
    parser.add_argument('--min-chars', type=int, default=200, dest='min_chars',
                        help='scarta file con meno di N caratteri dopo pulizia (default: 200)')

    args = parser.parse_args()

    if not args.inputs and not args.dir:
        print("  ERRORE: specifica almeno un file o --dir")
        parser.print_help()
        sys.exit(1)

    paths = collect_paths(args)
    if not paths:
        print("  ERRORE: nessun file trovato con i criteri specificati")
        sys.exit(1)

    out_path = Path(args.out)
    os.makedirs(out_path.parent, exist_ok=True)

    print(f"\n  convert_to_text.py — {len(paths)} file da convertire")
    print(f"  Output : {out_path}  ({'append' if args.append else 'overwrite'})\n")

    convert_files(paths, out_path, args.append, args.min_chars)


if __name__ == '__main__':
    main()
